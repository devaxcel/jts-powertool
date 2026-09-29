import csv
import io
import logging
import os
import re
import zipfile
from typing import List, Optional, Tuple
import httpx

logger = logging.getLogger(__name__)


def generate_pure_pdf(text: str, title: str = "Document") -> bytes:
    """Zero-dependency pure Python PDF 1.4 generator fallback."""
    lines = []
    if title:
        lines.append(title.upper())
        lines.append("=" * min(len(title), 60))
        lines.append("")
    for raw_line in text.strip().splitlines():
        # Wrap long lines roughly at 80 chars
        while len(raw_line) > 80:
            split_at = raw_line.rfind(" ", 0, 80)
            if split_at == -1:
                split_at = 80
            lines.append(raw_line[:split_at])
            raw_line = raw_line[split_at:].strip()
        lines.append(raw_line)

    # Simple multi-page text layout
    pages = []
    lines_per_page = 48
    for i in range(0, max(len(lines), 1), lines_per_page):
        pages.append(lines[i:i + lines_per_page])

    objects = []
    # 1: Catalog
    objects.append(b"1 0 obj <</Type /Catalog /Pages 2 0 R>> endobj")
    # 2: Pages container (kids updated later)
    # Font object will be obj 3
    objects.append(b"3 0 obj <</Type /Font /Subtype /Type1 /BaseFont /Helvetica>> endobj")

    page_obj_ids = []
    next_id = 4
    page_contents = []

    for page_lines in pages:
        page_id = next_id
        content_id = next_id + 1
        next_id += 2
        page_obj_ids.append(page_id)

        stream_ops = ["BT", "/F1 10 Tf", "50 740 Td", "14 TL"]
        for pl in page_lines:
            safe = pl.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            stream_ops.append(f"({safe}) '")
        stream_ops.append("ET")
        stream_bytes = "\n".join(stream_ops).encode("latin-1", errors="replace")

        page_obj = f"{page_id} 0 obj <</Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R /Resources <</Font <</F1 3 0 R>>>>>> endobj".encode("ascii")
        content_obj = f"{content_id} 0 obj <</Length {len(stream_bytes)}>> stream\n".encode("ascii") + stream_bytes + b"\nendstream\nendobj"
        page_contents.extend([page_obj, content_obj])

    kids_str = " ".join(f"{pid} 0 R" for pid in page_obj_ids)
    pages_obj = f"2 0 obj <</Type /Pages /Kids [{kids_str}] /Count {len(page_obj_ids)}>> endobj".encode("ascii")

    all_objs = [objects[0], pages_obj, objects[1]] + page_contents
    header = b"%PDF-1.4\n"
    body = b"\n".join(all_objs) + b"\n"

    # Xref table
    xref = [b"xref", f"0 {len(all_objs) + 1}".encode("ascii"), b"0000000000 65535 f "]
    offset = len(header)
    for o in all_objs:
        xref.append(f"{offset:010d} 00000 n ".encode("ascii"))
        offset += len(o) + 1

    trailer = f"trailer <</Size {len(all_objs) + 1} /Root 1 0 R>>\nstartxref\n{offset}\n%%EOF".encode("ascii")
    return header + body + b"\n".join(xref) + b"\n" + trailer


def generate_pdf(content: str, title: str = "Document") -> bytes:
    """Generate a clean styled PDF document using ReportLab (with pure Python fallback)."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib import colors

        buf = io.BytesIO()
        doc = SimpleDocTemplate(
            buf,
            pagesize=letter,
            leftMargin=50,
            rightMargin=50,
            topMargin=50,
            bottomMargin=50,
        )
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Title"],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#1A365D"),
            spaceAfter=12,
        )
        h1_style = ParagraphStyle(
            "H1",
            parent=styles["Heading1"],
            fontSize=13,
            leading=17,
            textColor=colors.HexColor("#2B6CB0"),
            spaceBefore=10,
            spaceAfter=5,
        )
        h2_style = ParagraphStyle(
            "H2",
            parent=styles["Heading2"],
            fontSize=11,
            leading=15,
            textColor=colors.HexColor("#2D3748"),
            spaceBefore=8,
            spaceAfter=4,
        )
        body_style = ParagraphStyle(
            "Body",
            parent=styles["Normal"],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#2D3748"),
            spaceAfter=5,
        )
        bullet_style = ParagraphStyle(
            "Bullet",
            parent=body_style,
            leftIndent=15,
            spaceAfter=3,
        )

        story = []
        if title:
            story.append(Paragraph(title, title_style))
            story.append(Spacer(1, 8))

        for line in content.strip().split("\n"):
            line_s = line.strip()
            if not line_s:
                story.append(Spacer(1, 5))
                continue

            if line_s.startswith("# "):
                h_text = line_s[2:].strip().replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(h_text, h1_style))
            elif line_s.startswith("## "):
                h_text = line_s[3:].strip().replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(h_text, h2_style))
            elif line_s.startswith("### "):
                h_text = line_s[4:].strip().replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(h_text, h2_style))
            elif line_s.startswith(("- ", "* ", "• ")):
                bullet_item = line_s[2:].strip().replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                bullet_item = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", bullet_item)
                story.append(Paragraph(f"&bull; {bullet_item}", bullet_style))
            else:
                safe_text = line_s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                safe_text = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", safe_text)
                story.append(Paragraph(safe_text, body_style))

        doc.build(story)
        return buf.getvalue()
    except Exception as e:
        logger.warning(f"ReportLab PDF generation failed, using pure Python fallback: {e}")
        return generate_pure_pdf(content, title)


def generate_pure_docx(text: str, title: str = "Document") -> bytes:
    """Zero-dependency pure Python OpenXML DOCX generator."""
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">\n'
            '  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>\n'
            '  <Default Extension="xml" ContentType="application/xml"/>\n'
            '  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>\n'
            "</Types>",
        )
        z.writestr(
            "_rels/.rels",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">\n'
            '  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>\n'
            "</Relationships>",
        )

        p_elements = []
        if title:
            safe_title = title.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            p_elements.append(
                f'<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:rPr><w:b/><w:sz w:val="36"/></w:rPr><w:t>{safe_title}</w:t></w:r></w:p>'
            )

        for line in text.strip().splitlines():
            safe_line = line.strip().replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            if not safe_line:
                p_elements.append("<w:p/>")
            elif safe_line.startswith("# "):
                p_elements.append(
                    f'<w:p><w:r><w:rPr><w:b/><w:sz w:val="28"/><w:color w:val="2B6CB0"/></w:rPr><w:t>{safe_line[2:]}</w:t></w:r></w:p>'
                )
            elif safe_line.startswith("## "):
                p_elements.append(
                    f'<w:p><w:r><w:rPr><w:b/><w:sz w:val="24"/><w:color w:val="2D3748"/></w:rPr><w:t>{safe_line[3:]}</w:t></w:r></w:p>'
                )
            elif safe_line.startswith(("- ", "* ", "• ")):
                p_elements.append(
                    f'<w:p><w:r><w:t>•  {safe_line[2:]}</w:t></w:r></w:p>'
                )
            else:
                p_elements.append(f"<w:p><w:r><w:t>{safe_line}</w:t></w:r></w:p>")

        doc_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">\n'
            f'  <w:body>{"".join(p_elements)}</w:body>\n'
            "</w:document>"
        )
        z.writestr("word/document.xml", doc_xml)
    return bio.getvalue()


def generate_docx(content: str, title: str = "Document") -> bytes:
    """Generate a Word (.docx) document using python-docx with pure OpenXML fallback."""
    try:
        import docx
        doc = docx.Document()
        if title:
            doc.add_heading(title, level=0)
        for line in content.strip().split("\n"):
            line_s = line.strip()
            if not line_s:
                continue
            if line_s.startswith("# "):
                doc.add_heading(line_s[2:].strip(), level=1)
            elif line_s.startswith("## "):
                doc.add_heading(line_s[3:].strip(), level=2)
            elif line_s.startswith("### "):
                doc.add_heading(line_s[4:].strip(), level=3)
            elif line_s.startswith(("- ", "* ", "• ")):
                doc.add_paragraph(line_s[2:].strip(), style="List Bullet")
            else:
                doc.add_paragraph(line_s)
        bio = io.BytesIO()
        doc.save(bio)
        return bio.getvalue()
    except Exception as e:
        logger.warning(f"python-docx generation failed, using pure OpenXML fallback: {e}")
        return generate_pure_docx(content, title)


def generate_xlsx(content: str, title: str = "Sheet1") -> bytes:
    """Zero-dependency pure Python OpenXML XLSX spreadsheet generator."""
    # Parse rows: CSV, TSV, pipe-delimited, or plain text lines
    rows: List[List[str]] = []
    lines = content.strip().splitlines()
    for line in lines:
        line_s = line.strip()
        if not line_s:
            continue
        if "|" in line_s:
            # Pipe-separated table row
            cells = [c.strip() for c in line_s.split("|") if c.strip() or line_s.startswith("|")]
            if cells and not all(set(c).issubset({"-", ":", " "}) for c in cells):
                rows.append(cells)
        elif "\t" in line_s:
            rows.append(line_s.split("\t"))
        elif "," in line_s:
            # Parse standard CSV line
            try:
                parsed = list(csv.reader([line_s]))[0]
                rows.append(parsed)
            except Exception:
                rows.append(line_s.split(","))
        else:
            rows.append([line_s])

    if not rows:
        rows = [["Data"]]

    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">\n'
            '  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>\n'
            '  <Default Extension="xml" ContentType="application/xml"/>\n'
            '  <Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>\n'
            '  <Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>\n'
            "</Types>",
        )
        z.writestr(
            "_rels/.rels",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">\n'
            '  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>\n'
            "</Relationships>",
        )
        z.writestr(
            "xl/_rels/workbook.xml.rels",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">\n'
            '  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>\n'
            "</Relationships>",
        )
        z.writestr(
            "xl/workbook.xml",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">\n'
            f'  <sheets><sheet name="{title[:30]}" sheetId="1" r:id="rId1"/></sheets>\n'
            "</workbook>",
        )

        sheet_data_lines = []
        for r_idx, row in enumerate(rows, 1):
            cell_strs = []
            for c_idx, val in enumerate(row):
                col_letter = chr(65 + c_idx) if c_idx < 26 else f"A{chr(65 + c_idx - 26)}"
                cell_ref = f"{col_letter}{r_idx}"
                s_val = str(val).strip().replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                cell_strs.append(f'<c r="{cell_ref}" t="inlineStr"><is><t>{s_val}</t></is></c>')
            sheet_data_lines.append(f'<row r="{r_idx}">{" ".join(cell_strs)}</row>')

        sheet_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">\n'
            f'  <sheetData>{"".join(sheet_data_lines)}</sheetData>\n'
            "</worksheet>"
        )
        z.writestr("xl/worksheets/sheet1.xml", sheet_xml)

    return bio.getvalue()


def generate_csv(content: str) -> bytes:
    """Generate CSV bytes from markdown table or text lines."""
    rows = []
    for line in content.strip().splitlines():
        line_s = line.strip()
        if not line_s:
            continue
        if "|" in line_s:
            cells = [c.strip() for c in line_s.split("|") if c.strip() or line_s.startswith("|")]
            if cells and not all(set(c).issubset({"-", ":", " "}) for c in cells):
                rows.append(cells)
        else:
            rows.append([line_s])

    buf = io.StringIO()
    writer = csv.writer(buf)
    for r in rows:
        writer.writerow(r)
    return buf.getvalue().encode("utf-8")


def generate_pptx(content: str, title: str = "Presentation") -> bytes:
    """Zero-dependency pure Python OpenXML PPTX presentation generator."""
    # Split content by '# ' or '---' to form slides
    slides = []
    current_slide = {"title": title, "bullets": []}
    for line in content.strip().splitlines():
        line_s = line.strip()
        if not line_s:
            continue
        if line_s.startswith("# "):
            if current_slide["bullets"] or current_slide["title"] != title:
                slides.append(current_slide)
            current_slide = {"title": line_s[2:].strip(), "bullets": []}
        elif line_s == "---":
            if current_slide["bullets"] or current_slide["title"]:
                slides.append(current_slide)
            current_slide = {"title": "Slide", "bullets": []}
        elif line_s.startswith(("- ", "* ", "• ")):
            current_slide["bullets"].append(line_s[2:].strip())
        else:
            current_slide["bullets"].append(line_s)
    if current_slide["bullets"] or current_slide["title"]:
        slides.append(current_slide)

    if not slides:
        slides = [{"title": title, "bullets": ["Content"]}]

    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w", compression=zipfile.ZIP_DEFLATED) as z:
        content_types = [
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">\n'
            '  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>\n'
            '  <Default Extension="xml" ContentType="application/xml"/>\n'
            '  <Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>'
        ]
        for i in range(1, len(slides) + 1):
            content_types.append(
                f'  <Override PartName="/ppt/slides/slide{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
            )
        content_types.append("</Types>")
        z.writestr("[Content_Types].xml", "\n".join(content_types))

        z.writestr(
            "_rels/.rels",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">\n'
            '  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/>\n'
            "</Relationships>",
        )

        pres_rels = [
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        ]
        slide_ids = []
        for i in range(1, len(slides) + 1):
            pres_rels.append(
                f'  <Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide{i}.xml"/>'
            )
            slide_ids.append(f'<p:sldId id="{255 + i}" r:id="rId{i}"/>')
        pres_rels.append("</Relationships>")
        z.writestr("ppt/_rels/presentation.xml.rels", "\n".join(pres_rels))

        z.writestr(
            "ppt/presentation.xml",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">\n'
            f'  <p:sldIdLst>{"".join(slide_ids)}</p:sldIdLst>\n'
            "</p:presentation>",
        )

        for i, s in enumerate(slides, 1):
            safe_title = s["title"].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            bullet_xmls = []
            for b in s["bullets"]:
                safe_b = b.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                bullet_xmls.append(f"<a:p><a:r><a:t>{safe_b}</a:t></a:r></a:p>")

            slide_xml = (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">\n'
                "  <p:cSld><p:spTree>\n"
                f'    <p:sp><p:txBody><a:p><a:r><a:t>{safe_title}</a:t></a:r></a:p></p:txBody></p:sp>\n'
                f'    <p:sp><p:txBody>{"".join(bullet_xmls)}</p:txBody></p:sp>\n'
                "  </p:spTree></p:cSld>\n"
                "</p:sld>"
            )
            z.writestr(f"ppt/slides/slide{i}.xml", slide_xml)

    return bio.getvalue()


def generate_file_bytes(filename: str, content: str) -> Tuple[bytes, str]:
    """
    Generate file bytes based on the target filename extension.
    Returns (bytes, mime_type).
    """
    fn_lower = filename.lower()

    if fn_lower.endswith(".pdf"):
        title = os.path.splitext(filename)[0].replace("_", " ").title()
        return generate_pdf(content, title), "application/pdf"

    if fn_lower.endswith(".docx"):
        title = os.path.splitext(filename)[0].replace("_", " ").title()
        return generate_docx(content, title), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    if fn_lower.endswith(".xlsx"):
        title = os.path.splitext(filename)[0].replace("_", " ").title()
        return generate_xlsx(content, title), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    if fn_lower.endswith(".csv"):
        return generate_csv(content), "text/csv"

    if fn_lower.endswith(".pptx"):
        title = os.path.splitext(filename)[0].replace("_", " ").title()
        return generate_pptx(content, title), "application/vnd.openxmlformats-officedocument.presentationml.presentation"

    if fn_lower.endswith(".html") or fn_lower.endswith(".htm"):
        html_content = content
        if not "<html" in content.lower():
            html_content = f"<!DOCTYPE html><html><head><meta charset='utf-8'><title>{filename}</title></head><body><pre>{content}</pre></body></html>"
        return html_content.encode("utf-8"), "text/html"

    if fn_lower.endswith(".json"):
        return content.encode("utf-8"), "application/json"

    # All text / markdown / code / scripts / other formats
    return content.encode("utf-8"), "text/plain"


def extract_file_generation_requests(text: str) -> List[Tuple[str, str]]:
    """
    Finds file generation blocks of the form:
    ```generate_file:filename.ext
    content
    ```
    Returns a list of (filename, content) tuples.
    """
    pattern = r"```(?:generate_file:)?([a-zA-Z0-9_\-\.]+\.(?:pdf|docx|xlsx|csv|pptx|txt|md|py|js|ts|html|json|xml|yaml|yml|sql|sh))\s*\n([\s\S]*?)```"
    matches = re.findall(pattern, text)
    results = []
    for fn, content in matches:
        fn_clean = fn.strip()
        if fn_clean and content.strip():
            results.append((fn_clean, content.strip()))
    return results


def strip_file_generation_blocks(text: str) -> str:
    """Replaces raw file generation blocks with clean presentation text."""
    pattern = r"```(?:generate_file:)?([a-zA-Z0-9_\-\.]+\.(?:pdf|docx|xlsx|csv|pptx|txt|md|py|js|ts|html|json|xml|yaml|yml|sql|sh))\s*\n[\s\S]*?```"

    def _replace(match):
        fn = match.group(1).strip()
        return f"\n📎 *Generated file:* `{fn}` (attached below)\n"

    cleaned = re.sub(pattern, _replace, text)
    return cleaned.strip()


def format_text_for_slack(text: str) -> str:
    """
    Converts standard markdown formatting to Slack mrkdwn format:
    - Converts **bold text** to *bold text* without altering code inside code blocks (```) or inline code (`...`).
    - Converts ***bold italic*** to *_bold italic_*.
    """
    if not text:
        return ""

    # Split text by code blocks (```...```) so code block contents are preserved
    parts = re.split(r"(```[\s\S]*?```)", text)
    converted_parts = []
    for part in parts:
        if part.startswith("```") and part.endswith("```"):
            converted_parts.append(part)
        else:
            # Protect inline code (`...`) from asterisk conversion
            sub_parts = re.split(r"(`[^`\n]+`)", part)
            sub_converted = []
            for sub_part in sub_parts:
                if sub_part.startswith("`") and sub_part.endswith("`"):
                    sub_converted.append(sub_part)
                else:
                    # Convert ***text*** to *_text_*
                    p = re.sub(r"\*\*\*(.*?)\*\*\*", r"*_\1_*", sub_part)
                    # Convert **text** to *text*
                    p = re.sub(r"\*\*(.*?)\*\*", r"*\1*", p)
                    sub_converted.append(p)
            converted_parts.append("".join(sub_converted))

    return "".join(converted_parts)


async def upload_file_to_slack(
    *,
    token: str,
    channel_id: str,
    thread_ts: Optional[str] = None,
    filename: str,
    file_bytes: bytes,
    title: Optional[str] = None,
    initial_comment: Optional[str] = None,
) -> bool:
    """
    Uploads a file to Slack using modern external upload API (files.getUploadURLExternal + completeUploadExternal),
    with automatic fallback to legacy files.upload.
    """
    if not token or not channel_id or not file_bytes:
        logger.error(f"Missing required parameters for Slack file upload: token={bool(token)}, channel={channel_id}, bytes={len(file_bytes)}")
        return False

    display_title = title or filename

    async with httpx.AsyncClient(timeout=60.0, follow_redirects=True) as client:
        # Method 1: Modern Slack API (getUploadURLExternal -> POST bytes -> completeUploadExternal)
        try:
            get_url_resp = await client.post(
                "https://slack.com/api/files.getUploadURLExternal",
                headers={"Authorization": f"Bearer {token}"},
                params={"filename": filename, "length": len(file_bytes)},
            )
            data = get_url_resp.json()
            if data.get("ok"):
                upload_url = data.get("upload_url")
                file_id = data.get("file_id")

                upload_resp = await client.post(upload_url, content=file_bytes)
                if upload_resp.status_code in [200, 201]:
                    complete_payload = {
                        "files": [{"id": file_id, "title": display_title}],
                        "channel_id": channel_id,
                    }
                    if thread_ts and not thread_ts.startswith("dm_"):
                        complete_payload["thread_ts"] = thread_ts
                    if initial_comment:
                        complete_payload["initial_comment"] = initial_comment

                    complete_resp = await client.post(
                        "https://slack.com/api/files.completeUploadExternal",
                        headers={
                            "Authorization": f"Bearer {token}",
                            "Content-Type": "application/json; charset=utf-8",
                        },
                        json=complete_payload,
                    )
                    complete_data = complete_resp.json()
                    if complete_data.get("ok"):
                        logger.info(f"Successfully uploaded '{filename}' to Slack channel {channel_id} via external upload API.")
                        return True
                    else:
                        logger.warning(f"files.completeUploadExternal error: {complete_data.get('error')}")
            else:
                logger.warning(f"files.getUploadURLExternal error: {data.get('error')}")
        except Exception as e:
            logger.warning(f"Slack external upload failed, trying legacy files.upload fallback: {e}")

        # Method 2: Fallback to classic files.upload
        try:
            files_multipart = {"file": (filename, file_bytes)}
            data_multipart = {
                "channels": channel_id,
                "filename": filename,
                "title": display_title,
            }
            if thread_ts and not thread_ts.startswith("dm_"):
                data_multipart["thread_ts"] = thread_ts
            if initial_comment:
                data_multipart["initial_comment"] = initial_comment

            legacy_resp = await client.post(
                "https://slack.com/api/files.upload",
                headers={"Authorization": f"Bearer {token}"},
                files=files_multipart,
                data=data_multipart,
            )
            legacy_data = legacy_resp.json()
            if legacy_data.get("ok"):
                logger.info(f"Successfully uploaded '{filename}' to Slack via legacy files.upload.")
                return True
            else:
                logger.error(f"Legacy files.upload error: {legacy_data.get('error')}")
                return False
        except Exception as e:
            logger.error(f"Legacy files.upload exception for '{filename}': {e}", exc_info=True)
            return False
